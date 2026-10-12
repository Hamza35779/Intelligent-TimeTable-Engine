import csv
import html
import io
import json
import textwrap
from datetime import datetime, timezone
from typing import List, Dict, Any, Tuple
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import landscape, A4
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.lib import colors
from PIL import Image, ImageDraw, ImageFont
from utils.timetable_formatter import (
    build_table_rows,
    build_university_sheet_data,
    VIEW_TITLES,
)


ALLOWED_EXPORT_VIEWS = ("master", "teacher", "student", "staff", "canteen")
DANGEROUS_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def _sanitize_spreadsheet_cell(value: Any) -> str:
    text = str(value if value is not None else "")
    stripped = text.lstrip()
    if stripped.startswith(DANGEROUS_FORMULA_PREFIXES):
        return f"'{text}"
    return text


def _wrap_and_truncate_cell_text(text: str, max_chars_per_line: int = 22, max_lines: int = 3) -> List[str]:
    """
    Wrap text to fit within cell width and truncate with an ellipsis (...) if it exceeds max_lines.
    Prevents any text overflow in PNG/JPG image renders (Step 6).
    """
    raw_lines = [ln.strip() for ln in str(text or "").split("\n") if ln.strip()]
    wrapped: List[str] = []
    for raw in raw_lines:
        chunks = textwrap.wrap(raw, width=max_chars_per_line) or [raw[:max_chars_per_line]]
        wrapped.extend(chunks)

    if len(wrapped) > max_lines:
        kept = wrapped[:max_lines]
        last = kept[-1]
        if len(last) > max_chars_per_line - 3:
            kept[-1] = last[: max_chars_per_line - 3] + "..."
        else:
            kept[-1] = last + "..."
        return kept
    return wrapped


class TimetableExporter:
    def __init__(
        self,
        entries: List[Dict[str, Any]],
        metadata: Dict[str, Any],
        view_type: str = "master",
    ):
        self.entries = entries
        self.metadata = metadata or {}
        v_clean = (view_type or "master").strip().lower()
        if v_clean not in ALLOWED_EXPORT_VIEWS:
            raise ValueError(
                f"Invalid view '{view_type}'. Allowed views: {list(ALLOWED_EXPORT_VIEWS)}"
            )
        self.view_type = v_clean
        self.schedule_id = self.metadata.get("schedule_id", 1)
        self.generated_at = (
            self.metadata.get("generated_at") or datetime.now(timezone.utc).isoformat()
        )
        self.section_title = self.metadata.get("section_title") or "AI 2024 - Section A"
        self.title = VIEW_TITLES.get(self.view_type, self.section_title)
        self.rows = build_table_rows(
            entries, view_type=self.view_type, section_title=self.section_title
        )
        self.sheet = build_university_sheet_data(
            entries, view_type=self.view_type, section_title=self.section_title
        )

    def export(self, format_type: str) -> Tuple[bytes, str, str]:
        fmt = format_type.strip().lower()
        ts = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
        base_name = f"timetable_{self.view_type}_{ts}"

        if fmt == "csv":
            return self._export_csv(), "text/csv", f"{base_name}.csv"
        elif fmt == "xlsx":
            return (
                self._export_xlsx(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                f"{base_name}.xlsx",
            )
        elif fmt == "pdf":
            return self._export_pdf(), "application/pdf", f"{base_name}.pdf"
        elif fmt in ("png", "jpg", "jpeg"):
            ext = "jpg" if fmt in ("jpg", "jpeg") else "png"
            mime = "image/jpeg" if ext == "jpg" else "image/png"
            return self._export_image(ext), mime, f"{base_name}.{ext}"
        elif fmt == "json":
            return self._export_json(), "application/json", f"{base_name}.json"
        elif fmt == "html":
            return self._export_html(), "text/html", f"{base_name}.html"
        else:
            raise ValueError(f"Unsupported export format: {format_type}")

    def _export_csv(self) -> bytes:
        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow(
            [
                _sanitize_spreadsheet_cell(self.title),
                f"Schedule ID: {self.schedule_id}",
                f"Generated At: {self.generated_at}",
            ]
        )
        sanitized_rows = [
            [_sanitize_spreadsheet_cell(cell) for cell in row] for row in self.rows
        ]
        writer.writerows(sanitized_rows)
        writer.writerow([])
        writer.writerow(
            [
                "Serial",
                "Course Code",
                "Course Titles",
                "CHs",
                "Name of Faculty",
                "Elective / Compulsory",
                "Venues",
                "Remarks",
            ]
        )
        for c_row in self.sheet["course_legend"]:
            writer.writerow(
                [
                    _sanitize_spreadsheet_cell(c_row["serial"]),
                    _sanitize_spreadsheet_cell(c_row["course_code"]),
                    _sanitize_spreadsheet_cell(c_row["course_title"]),
                    _sanitize_spreadsheet_cell(c_row["credit_hours"]),
                    _sanitize_spreadsheet_cell(c_row["faculty_name"]),
                    _sanitize_spreadsheet_cell(c_row["elective_compulsory"]),
                    _sanitize_spreadsheet_cell(c_row["venues"]),
                    _sanitize_spreadsheet_cell(c_row["remarks"]),
                ]
            )
        return buf.getvalue().encode("utf-8")

    def _populate_view_worksheet(self, ws, v_type: str):
        black_border = Border(
            left=Side(style="thin", color="000000"),
            right=Side(style="thin", color="000000"),
            top=Side(style="thin", color="000000"),
            bottom=Side(style="thin", color="000000"),
        )
        blue_fill = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
        center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

        v_rows = build_table_rows(
            self.entries, view_type=v_type, section_title=self.section_title
        )
        v_sheet = build_university_sheet_data(
            self.entries, view_type=v_type, section_title=self.section_title
        )
        v_title = f"{VIEW_TITLES.get(v_type, self.section_title)} (Schedule #{self.schedule_id})"

        num_cols = max(10, len(v_rows[0]) if v_rows else 6)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=num_cols)
        for col_idx in range(1, num_cols + 1):
            c = ws.cell(row=1, column=col_idx)
            c.fill = blue_fill
            c.border = black_border
        banner = ws.cell(row=1, column=1, value=_sanitize_spreadsheet_cell(v_title))
        banner.font = Font(name="Calibri", color="FFFFFF", bold=True, size=11)
        banner.alignment = center_align
        ws.row_dimensions[1].height = 24

        current_row = 2
        for r_idx, row in enumerate(v_rows):
            for c_idx, val in enumerate(row, start=1):
                c = ws.cell(
                    row=current_row,
                    column=c_idx,
                    value=_sanitize_spreadsheet_cell(val),
                )
                c.border = black_border
                c.alignment = center_align
                c.font = Font(name="Calibri", bold=(r_idx == 0 or c_idx == 1), size=10)
            ws.row_dimensions[current_row].height = 24 if r_idx == 0 else 38
            current_row += 1

        current_row += 1
        # Append Course Legend table
        leg_headers = [
            "Serial",
            "Course Code",
            "Course Titles",
            "CHs",
            "Name of Faculty",
            "Elective / Compulsory",
            "Venues",
            "Remarks",
        ]
        for c_idx, h_val in enumerate(leg_headers, start=1):
            hc = ws.cell(row=current_row, column=c_idx, value=h_val)
            hc.font = Font(name="Calibri", bold=True, size=10)
            hc.border = black_border
            hc.alignment = center_align
        current_row += 1

        for c_row in v_sheet["course_legend"]:
            vals = [
                c_row["serial"],
                c_row["course_code"],
                c_row["course_title"],
                c_row["credit_hours"],
                c_row["faculty_name"],
                c_row["elective_compulsory"],
                c_row["venues"],
                c_row["remarks"],
            ]
            for c_idx, val in enumerate(vals, start=1):
                vc = ws.cell(
                    row=current_row,
                    column=c_idx,
                    value=_sanitize_spreadsheet_cell(val),
                )
                vc.font = Font(name="Calibri", size=10)
                vc.border = black_border
                vc.alignment = center_align
            current_row += 1

        for col_idx in range(1, num_cols + 1):
            ws.column_dimensions[get_column_letter(col_idx)].width = 20

    def _export_xlsx(self) -> bytes:
        """
        Step 6 XLSX rule:
        One sheet per view (with the requested view placed first), plus a Metrics sheet.
        """
        wb = Workbook()
        ordered_views = [self.view_type] + [
            v for v in ALLOWED_EXPORT_VIEWS if v != self.view_type
        ]

        for idx, v_type in enumerate(ordered_views):
            sheet_title = f"{v_type.capitalize()} View"
            ws = wb.active if idx == 0 else wb.create_sheet(sheet_title)
            ws.title = sheet_title
            self._populate_view_worksheet(ws, v_type)

        ws_meta = wb.create_sheet("Metrics")
        ws_meta.append(["Metric", "Value"])
        ws_meta.append(["Schedule ID", self.schedule_id])
        ws_meta.append(["Generated Timestamp", self.generated_at])
        ws_meta.append(["Primary Requested View", self.title])
        ws_meta.append(["Fairness Score", self.metadata.get("fairness_score", 0.92)])
        ws_meta.append(
            ["Preference Satisfaction", self.metadata.get("preference_satisfaction", 0.88)]
        )
        ws_meta.append(
            ["Total Scheduled Hours", self.metadata.get("total_courses", len(self.entries))]
        )
        ws_meta.append(["Total Teachers", self.metadata.get("total_teachers", 6)])

        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()

    def _export_pdf(self) -> bytes:
        """
        Step 6 PDF rule:
        Landscape A4, header with schedule ID, view title, and timestamp, rendering
        the requested view's table and course summary legend.
        """
        buf = io.BytesIO()
        doc = SimpleDocTemplate(
            buf,
            pagesize=landscape(A4),
            leftMargin=20,
            rightMargin=20,
            topMargin=20,
            bottomMargin=20,
        )
        elements = []
        styles = getSampleStyleSheet()
        header_style = ParagraphStyle(
            "PDFHeader",
            parent=styles["Heading2"],
            fontName="Helvetica-Bold",
            fontSize=12,
            textColor=colors.HexColor("#1F4E78"),
            alignment=1,
            spaceAfter=4,
        )
        sub_style = ParagraphStyle(
            "PDFSubHeader",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=8.5,
            textColor=colors.HexColor("#334155"),
            alignment=1,
            spaceAfter=10,
        )
        cell_style = ParagraphStyle(
            "CellText",
            fontName="Helvetica",
            fontSize=7.5,
            leading=9.5,
            alignment=1,
        )
        bold_style = ParagraphStyle(
            "CellBold",
            fontName="Helvetica-Bold",
            fontSize=8,
            leading=10,
            alignment=1,
        )

        elements.append(
            Paragraph(
                f"{html.escape(self.title)} — Schedule #{self.schedule_id}",
                header_style,
            )
        )
        elements.append(
            Paragraph(
                f"Generated At: {html.escape(str(self.generated_at))} | View: {html.escape(self.view_type.upper())}",
                sub_style,
            )
        )

        # Render the requested view's table rows
        num_cols = len(self.rows[0]) if self.rows else 6
        avail_w = 10.6 * inch
        col_widths = [avail_w / num_cols] * num_cols

        pdf_table_data = []
        for r_idx, row in enumerate(self.rows):
            st = bold_style if r_idx == 0 else cell_style
            pdf_table_data.append(
                [
                    Paragraph(html.escape(str(c)).replace("\n", "<br/>"), st)
                    for c in row
                ]
            )

        view_table = Table(pdf_table_data, colWidths=col_widths, repeatRows=1)
        view_table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2E75B6")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("GRID", (0, 0), (-1, -1), 0.6, colors.black),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
                ]
            )
        )
        elements.append(view_table)
        elements.append(Spacer(1, 12))

        # Course Summary Legend
        legend_headers = [
            "Serial",
            "Course Code",
            "Course Titles",
            "CHs",
            "Name of Faculty",
            "Elective / Compulsory",
            "Venues",
            "Remarks",
        ]
        lower_data = [[Paragraph(html.escape(h), bold_style) for h in legend_headers]]
        for c_row in self.sheet["course_legend"]:
            lower_data.append(
                [
                    Paragraph(html.escape(str(c_row["serial"])), cell_style),
                    Paragraph(html.escape(str(c_row["course_code"])), cell_style),
                    Paragraph(html.escape(str(c_row["course_title"])), cell_style),
                    Paragraph(html.escape(str(c_row["credit_hours"])), cell_style),
                    Paragraph(html.escape(str(c_row["faculty_name"])), cell_style),
                    Paragraph(html.escape(str(c_row["elective_compulsory"])), cell_style),
                    Paragraph(html.escape(str(c_row["venues"])), cell_style),
                    Paragraph(html.escape(str(c_row["remarks"])), cell_style),
                ]
            )
        lower_col_widths = [
            0.6 * inch,
            1.0 * inch,
            2.5 * inch,
            0.6 * inch,
            2.0 * inch,
            1.5 * inch,
            1.2 * inch,
            1.19 * inch,
        ]
        lower_table = Table(lower_data, colWidths=lower_col_widths)
        lower_table.setStyle(
            TableStyle(
                [
                    ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("GRID", (0, 0), (-1, -1), 0.6, colors.black),
                    ("BACKGROUND", (0, 0), (-1, -1), colors.white),
                ]
            )
        )
        elements.append(lower_table)

        doc.build(elements)
        return buf.getvalue()

    def _export_image(self, ext: str = "png") -> bytes:
        """
        Step 6 PNG/JPG rule:
        Text must not overflow cells. Wrap or truncate with an ellipsis (...).
        """
        margin = 20
        num_cols = max(1, len(self.rows[0]) if self.rows else 6)
        col_w = 175
        total_w = col_w * num_cols
        banner_h = 38
        row_h = 62

        num_rows = len(self.rows)
        img_w = total_w + margin * 2
        img_h = margin * 2 + banner_h + num_rows * row_h

        img = Image.new("RGB", (img_w, img_h), color=(255, 255, 255))
        draw = ImageDraw.Draw(img)
        font = ImageFont.load_default()

        def draw_fitted_cell(box: Tuple[int, int, int, int], text: str, fill=(0, 0, 0)):
            x0, y0, x1, y1 = box
            cell_width_px = max(40, x1 - x0 - 10)
            max_chars = max(10, cell_width_px // 6)
            lines = _wrap_and_truncate_cell_text(text, max_chars_per_line=max_chars, max_lines=3)
            if not lines:
                return
            line_h = 14
            total_txt_h = len(lines) * line_h
            start_y = y0 + max(4, (y1 - y0 - total_txt_h) // 2)
            for idx, line in enumerate(lines):
                approx_w = min(cell_width_px, len(line) * 6)
                start_x = x0 + max(5, (x1 - x0 - approx_w) // 2)
                draw.text((start_x, start_y + idx * line_h), line, fill=fill, font=font)

        y = margin
        # Header Banner with View Title, Schedule ID & Timestamp
        draw.rectangle(
            [(margin, y), (margin + total_w, y + banner_h)],
            fill=(46, 117, 182),
            outline=(0, 0, 0),
            width=1,
        )
        banner_txt = f"{self.title} | Schedule #{self.schedule_id} | {self.generated_at[:19]}"
        draw_fitted_cell((margin, y, margin + total_w, y + banner_h), banner_txt, fill=(255, 255, 255))
        y += banner_h

        for r_idx, row in enumerate(self.rows):
            bg = (235, 243, 250) if r_idx == 0 else (255, 255, 255)
            for c_idx, cell_val in enumerate(row):
                x0 = margin + c_idx * col_w
                draw.rectangle(
                    [(x0, y), (x0 + col_w, y + row_h)],
                    fill=bg,
                    outline=(0, 0, 0),
                    width=1,
                )
                draw_fitted_cell((x0, y, x0 + col_w, y + row_h), str(cell_val))
            y += row_h

        buf = io.BytesIO()
        if ext == "jpg":
            img.save(buf, format="JPEG", quality=95)
        else:
            img.save(buf, format="PNG")
        return buf.getvalue()

    def _export_json(self) -> bytes:
        payload = {
            "schedule_id": self.schedule_id,
            "generated_at": self.generated_at,
            "view_type": self.view_type,
            "title": self.title,
            "section_title": self.section_title,
            "metadata": {
                "schedule_id": self.schedule_id,
                "generated_at": self.generated_at,
                "fairness_score": self.metadata.get("fairness_score", 0.92),
                "preference_satisfaction": self.metadata.get("preference_satisfaction", 0.88),
                "total_courses": self.metadata.get("total_courses", len(self.entries)),
                "total_teachers": self.metadata.get("total_teachers", 6),
            },
            "headers": self.rows[0] if self.rows else [],
            "table_rows": self.rows[1:] if len(self.rows) > 1 else [],
            "sheet": self.sheet,
            "entries": self.entries,
        }
        return json.dumps(payload, indent=2).encode("utf-8")

    def _export_html(self) -> bytes:
        safe_title = html.escape(self.title)
        safe_ts = html.escape(str(self.generated_at))
        html_lines = [
            "<!DOCTYPE html>",
            "<html><head><meta charset='UTF-8'>",
            f"<title>{safe_title} - Schedule #{self.schedule_id}</title>",
            "<style>",
            "body { font-family: Calibri, 'Segoe UI', Arial, sans-serif; margin: 24px; background: #ffffff; color: #000000; }",
            "table { border-collapse: collapse; width: 100%; background: #ffffff; margin-bottom: 24px; }",
            "th, td { border: 1px solid #000000; padding: 8px 6px; text-align: center; vertical-align: middle; font-size: 13px; white-space: pre-line; }",
            ".banner { background: #2E75B6; color: #ffffff; font-weight: 700; font-size: 15px; padding: 8px; }",
            ".bold-cell { font-weight: 700; background: #f1f5f9; }",
            "</style></head><body>",
            "<table>",
            f"<thead><tr><th colspan='{len(self.rows[0]) if self.rows else 6}' class='banner'>{safe_title} (Schedule #{self.schedule_id} | {safe_ts})</th></tr>",
        ]

        for r_idx, row in enumerate(self.rows):
            if r_idx == 0:
                html_lines.append(
                    "<tr>"
                    + "".join(f"<th class='bold-cell'>{html.escape(str(c))}</th>" for c in row)
                    + "</tr></thead><tbody>"
                )
            else:
                html_lines.append(
                    "<tr>"
                    + "".join(f"<td>{html.escape(str(c))}</td>" for c in row)
                    + "</tr>"
                )

        html_lines.append("</tbody></table>")

        # Lower Table: Course & Faculty Summary Legend
        html_lines.extend(
            [
                "<table><thead><tr>",
                "<th>Serial</th><th>Course Code</th><th>Course Titles</th><th>CHs</th>",
                "<th>Name of Faculty</th><th>Elective / Compulsory</th><th>Venues</th><th>Remarks</th>",
                "</tr></thead><tbody>",
            ]
        )
        for c_row in self.sheet["course_legend"]:
            html_lines.append(
                "<tr>"
                f"<td>{html.escape(str(c_row['serial']))}</td>"
                f"<td>{html.escape(str(c_row['course_code']))}</td>"
                f"<td>{html.escape(str(c_row['course_title']))}</td>"
                f"<td>{html.escape(str(c_row['credit_hours']))}</td>"
                f"<td>{html.escape(str(c_row['faculty_name']))}</td>"
                f"<td>{html.escape(str(c_row['elective_compulsory']))}</td>"
                f"<td>{html.escape(str(c_row['venues']))}</td>"
                f"<td>{html.escape(str(c_row['remarks']))}</td>"
                "</tr>"
            )
        html_lines.append("</tbody></table></body></html>")
        return "\n".join(html_lines).encode("utf-8")


def export_to_csv(
    entries: List[Dict[str, Any]],
    view_type: str = "master",
    metadata: Dict[str, Any] = None,
) -> bytes:
    exporter = TimetableExporter(entries=entries, metadata=metadata or {}, view_type=view_type)
    return exporter.export("csv")[0]


def export_to_excel(
    entries: List[Dict[str, Any]],
    metadata: Dict[str, Any] = None,
    view_type: str = "master",
) -> bytes:
    exporter = TimetableExporter(entries=entries, metadata=metadata or {}, view_type=view_type)
    return exporter.export("xlsx")[0]


def export_to_pdf(
    entries: List[Dict[str, Any]],
    metadata: Dict[str, Any] = None,
    view_type: str = "master",
) -> bytes:
    exporter = TimetableExporter(entries=entries, metadata=metadata or {}, view_type=view_type)
    return exporter.export("pdf")[0]


def export_to_image(
    entries: List[Dict[str, Any]],
    img_format: str = "PNG",
    view_type: str = "master",
    metadata: Dict[str, Any] = None,
) -> bytes:
    fmt = "jpg" if img_format.upper() in ("JPG", "JPEG") else "png"
    exporter = TimetableExporter(entries=entries, metadata=metadata or {}, view_type=view_type)
    return exporter.export(fmt)[0]


def export_to_json(
    entries: List[Dict[str, Any]],
    metadata: Dict[str, Any] = None,
    view_type: str = "master",
) -> bytes:
    exporter = TimetableExporter(entries=entries, metadata=metadata or {}, view_type=view_type)
    return exporter.export("json")[0]


def export_to_html(
    entries: List[Dict[str, Any]],
    metadata: Dict[str, Any] = None,
    view_type: str = "master",
) -> bytes:
    exporter = TimetableExporter(entries=entries, metadata=metadata or {}, view_type=view_type)
    return exporter.export("html")[0]

