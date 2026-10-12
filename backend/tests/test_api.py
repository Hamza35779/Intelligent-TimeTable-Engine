import pytest
from fastapi.testclient import TestClient
from app.main import app
from app.database import Base, engine

client = TestClient(app)


@pytest.fixture(autouse=True)
def setup_database():
    Base.metadata.create_all(bind=engine)
    yield


def test_health_endpoint():
    response = client.get("/health")
    assert response.status_code == 200
    data = response.json()
    assert data["status"] in ("healthy", "degraded")
    assert "solver" in data
    assert data["solver"]["ready"] is True


def test_constraint_crud():
    payload = {
        "type": "soft",
        "name": "Maximize Morning Slots",
        "description": "Prefers scheduling before 12pm",
        "weight": 2.5,
        "constraint_data": {"prefer_morning": True},
    }
    create_resp = client.post("/constraints", json=payload)
    assert create_resp.status_code == 201
    created_data = create_resp.json()
    constraint_id = created_data["id"]

    list_resp = client.get("/constraints")
    assert list_resp.status_code == 200
    assert any(c["id"] == constraint_id for c in list_resp.json())

    del_resp = client.delete(f"/constraints/{constraint_id}")
    assert del_resp.status_code == 204


def test_schedule_generation_views_exports_feedback_and_chat():
    generate_payload = {
        "faculty": [
            {
                "name": "Ahmad Ali",
                "department": "Computer Science",
                "max_hours_per_day": 4,
                "max_days_per_week": 5,
                "email": "ahmad@univ.edu",
                "preferences": {"prefer_morning": True},
                "unavailable_slots": [],
            },
            {
                "name": "Fatima Khan",
                "department": "Information Theory",
                "max_hours_per_day": 4,
                "max_days_per_week": 5,
                "email": "fatima@univ.edu",
                "preferences": {"prefer_morning": False},
                "unavailable_slots": [],
            },
        ],
        "venues": [
            {"name": "Room 101", "capacity": 100, "equipment": ["projector"], "available_slots": [0, 1, 2, 3]},
            {"name": "Room 102", "capacity": 60, "equipment": ["computers"], "available_slots": [0, 1, 2, 3]},
        ],
        "slots": [
            {"day_of_week": "Monday", "start_time": "09:00", "end_time": "10:00", "duration_minutes": 60},
            {"day_of_week": "Monday", "start_time": "10:00", "end_time": "11:00", "duration_minutes": 60},
            {"day_of_week": "Tuesday", "start_time": "09:00", "end_time": "10:00", "duration_minutes": 60},
            {"day_of_week": "Tuesday", "start_time": "10:00", "end_time": "11:00", "duration_minutes": 60},
        ],
        "courses": [
            {"name": "Data Structures", "code": "CS201", "faculty_id": 0, "capacity_required": 50, "duration_minutes": 60},
            {"name": "Database Systems", "code": "CS202", "faculty_id": 1, "capacity_required": 40, "duration_minutes": 60},
        ],
        "timeout_seconds": 5,
    }

    gen_resp = client.post("/schedule/generate", json=generate_payload)
    assert gen_resp.status_code == 201
    gen_data = gen_resp.json()
    schedule_id = gen_data["schedule_id"]
    assert gen_data["analytics"]["fairness_score"] >= 0.80

    # Tabular endpoint
    table_resp = client.get(f"/schedule/{schedule_id}/table?view=master")
    assert table_resp.status_code == 200
    assert table_resp.json()["format"] == "tabular"

    # Multi-view endpoint (all 5 perspectives)
    views_resp = client.get(f"/schedule/{schedule_id}/views")
    assert views_resp.status_code == 200
    views_data = views_resp.json()["views"]
    assert set(views_data.keys()) == {"master", "teacher", "student", "staff", "canteen"}

    # Export in all 7 formats
    for fmt in ("csv", "xlsx", "pdf", "png", "jpg", "json", "html"):
        exp_resp = client.get(f"/export/{schedule_id}?format={fmt}&view=canteen")
        assert exp_resp.status_code == 200
        assert len(exp_resp.content) > 0

    # Submit Feedback -> Triggers ML Retraining
    fb_resp = client.post(
        f"/feedback/{schedule_id}",
        json={
            "teacher_id": 1,
            "teacher_name": "Ahmad Ali",
            "satisfaction_score": 5,
            "preferred_time": "Morning",
            "workload_satisfaction": 4,
            "comments": "Great morning allocation!",
        },
    )
    assert fb_resp.status_code == 201
    assert fb_resp.json()["received"] is True

    # ML stats & retrain endpoints
    ml_stats = client.get("/ml/model-stats")
    assert ml_stats.status_code == 200
    assert ml_stats.json()["training_samples"] >= 80

    ml_retrain = client.post("/ml/retrain")
    assert ml_retrain.status_code == 200

    # Conversational AI Chatbot workflow
    chat_start = client.post("/chat/start", json={"schedule_id": schedule_id})
    assert chat_start.status_code == 200
    conv_id = chat_start.json()["conversation_id"]

    chat_msg = client.post(
        f"/chat/{conv_id}/message",
        json={"message": "I need schedules for teachers, students, and canteen staff in PDF and Excel"},
    )
    assert chat_msg.status_code == 200
    msg_data = chat_msg.json()
    assert "teacher" in msg_data["detected_views"]
    assert "canteen" in msg_data["detected_views"]
    assert msg_data["status"] == "ready_to_generate"

    chat_gen = client.post(
        f"/chat/{conv_id}/generate",
        json={"views": ["master", "teacher", "student", "staff", "canteen"], "formats": ["pdf", "xlsx", "png"]},
    )
    assert chat_gen.status_code == 200
    assert chat_gen.json()["schedules_generated"] == 5


def test_csv_upload_endpoint():
    csv_content = (
        "teacher,course,code,hours,preferred_time,department,capacity\n"
        "Ahmad Ali,Data Structures,CS301,3,Morning,Computer Science,50\n"
        "Fatima Khan,Database Systems,CS302,4,Afternoon,Computer Science,45\n"
    )
    resp = client.post(
        "/schedule/upload-csv",
        files={"file": ("courses.csv", csv_content.encode("utf-8"), "text/csv")},
    )
    assert resp.status_code == 201
    # 3 credit hours + 4 credit hours = 7 contact hour slot entries
    assert len(resp.json()["schedule"]) == 7


def test_reschedule_scoping_and_unseated_slot_exclusion():
    """
    Regression test for [ISS-001]:
    Verify that rescheduling a specific run only reschedules that run's courses
    and strictly vacates the unseated slot without index shifting errors.
    """
    payload = {
        "faculty": [
            {
                "name": "Dr. Resched A",
                "department": "CS",
                "max_hours_per_day": 4,
                "max_days_per_week": 5,
                "email": "resched_a@univ.edu",
                "preferences": {"prefer_morning": True},
                "unavailable_slots": [],
            }
        ],
        "venues": [
            {"name": "Resched Hall 1", "capacity": 80, "equipment": [], "available_slots": [0, 1, 2]},
        ],
        "slots": [
            {"day_of_week": "Monday", "start_time": "09:00", "end_time": "10:00", "duration_minutes": 60},
            {"day_of_week": "Monday", "start_time": "10:00", "end_time": "11:00", "duration_minutes": 60},
            {"day_of_week": "Monday", "start_time": "11:00", "end_time": "12:00", "duration_minutes": 60},
        ],
        "courses": [
            {"name": "Systems I", "code": "SYS101", "faculty_id": 0, "capacity_required": 40, "duration_minutes": 60},
            {"name": "Systems II", "code": "SYS102", "faculty_id": 0, "capacity_required": 40, "duration_minutes": 60},
        ],
        "timeout_seconds": 5,
    }
    gen_resp = client.post("/schedule/generate", json=payload)
    assert gen_resp.status_code == 201
    initial_data = gen_resp.json()
    sched_id = initial_data["schedule_id"]
    assert len(initial_data["schedule"]) == 2

    # Unseat 0-based slot 0 (Monday 09:00)
    resched_resp = client.post(
        "/schedule/reschedule",
        json={
            "timetable_id": sched_id,
            "unseated_slot_ids": [0],
            "unavailable_venue_ids": [],
            "unavailable_faculty_slots": {},
            "timeout_seconds": 5,
        },
    )
    assert resched_resp.status_code == 200
    resched_data = resched_resp.json()
    # Must contain only the 2 courses from this schedule run, not courses from previous runs
    assert len(resched_data["schedule"]) == 2
    # Neither course may be scheduled at Monday 09:00 (slot index 0)
    for entry in resched_data["schedule"]:
        assert not (entry["day_of_week"] == "Monday" and entry["start_time"] == "09:00")


def test_exporter_xss_and_formula_injection_sanitization():
    """
    Regression test for [ISS-003]:
    Verify HTML export escapes script tags and CSV export neutralizes formula injection prefixes.
    """
    payload = {
        "faculty": [
            {
                "name": "=CMD|' /C calc'!A0",
                "department": "Security",
                "max_hours_per_day": 4,
                "max_days_per_week": 5,
                "email": "sec_audit@univ.edu",
                "preferences": {},
                "unavailable_slots": [],
            }
        ],
        "venues": [
            {"name": "Safe Room <img src=x onerror=alert(1)>", "capacity": 60, "equipment": []},
        ],
        "slots": [
            {"day_of_week": "Monday", "start_time": "09:00", "end_time": "10:00", "duration_minutes": 60},
        ],
        "courses": [
            {
                "name": "<script>alert('xss')</script>",
                "code": "SEC999",
                "faculty_id": 0,
                "capacity_required": 30,
                "duration_minutes": 60,
            },
        ],
        "timeout_seconds": 5,
    }
    gen_resp = client.post("/schedule/generate", json=payload)
    assert gen_resp.status_code == 201
    sched_id = gen_resp.json()["schedule_id"]

    html_resp = client.get(f"/export/{sched_id}?format=html&view=master")
    assert html_resp.status_code == 200
    html_text = html_resp.content.decode("utf-8")
    assert "<script>alert('xss')</script>" not in html_text
    assert "&lt;script&gt;alert(&#x27;xss&#x27;)&lt;/script&gt;" in html_text

    csv_resp = client.get(f"/export/{sched_id}?format=csv&view=master")
    assert csv_resp.status_code == 200
    csv_text = csv_resp.content.decode("utf-8")
    assert "'=CMD|' /C calc'!A0" in csv_text


def test_chatbot_keyword_fallback_and_invalid_rejection():
    """
    Step 9 Chatbot Test:
    - Keyword fallback parses "Give me teacher and student schedules in PDF and Excel"
      into views=["teacher","student"], formats=["pdf","xlsx"].
    - Invalid view or format is rejected (400 Bad Request).
    """
    msg_resp = client.post(
        "/chat/message",
        json={"message": "Give me teacher and student schedules in PDF and Excel"},
    )
    assert msg_resp.status_code == 200
    data = msg_resp.json()
    assert data["detected_views"] == ["teacher", "student"]
    assert data["detected_formats"] == ["pdf", "xlsx"]
    assert data["mode"] in ("keyword_fallback", "gemini")

    # Reject invalid view in /chat/message
    bad_msg = client.post(
        "/chat/message",
        json={"message": "Give me dean view in pdf"},
    )
    assert bad_msg.status_code == 400

    # Reject invalid format in /chat/generate-schedules
    bad_gen = client.post(
        "/chat/generate-schedules",
        json={"views": ["teacher"], "formats": ["docx"]},
    )
    assert bad_gen.status_code == 400


def test_all_7_exports_size_and_content_validation():
    """
    Step 9 Export Test:
    Each of the 7 formats produces a non-empty file > 100 bytes and containing
    at least one known course name from the schedule (for text formats: CSV, JSON, HTML;
    for XLSX, readable via openpyxl; for PDF/PNG/JPG, non-empty valid file header).
    """
    import io
    from openpyxl import load_workbook

    demo_resp = client.post("/demo/run")
    assert demo_resp.status_code == 201
    sched_id = demo_resp.json()["schedule_id"]
    known_course = "Artificial Intelligence"

    for fmt in ("csv", "xlsx", "pdf", "png", "jpg", "json", "html"):
        resp = client.get(f"/export/{sched_id}?format={fmt}&view=master")
        assert resp.status_code == 200
        raw = resp.content
        assert len(raw) > 100, f"Format {fmt} returned <= 100 bytes ({len(raw)})"

        if fmt in ("csv", "json", "html"):
            text = raw.decode("utf-8")
            assert known_course in text, f"Known course '{known_course}' missing in {fmt} export"
        elif fmt == "xlsx":
            wb = load_workbook(io.BytesIO(raw))
            assert "Metrics" in wb.sheetnames
            all_cells_text = " ".join(
                str(cell)
                for sheet in wb.worksheets
                for row in sheet.iter_rows(values_only=True)
                for cell in row
                if cell is not None
            )
            assert known_course in all_cells_text
        elif fmt == "pdf":
            assert raw.startswith(b"%PDF-")
        elif fmt == "png":
            assert raw.startswith(b"\x89PNG\r\n\x1a\n")
        elif fmt == "jpg":
            assert raw.startswith(b"\xff\xd8\xff")


def test_end_to_end_demo_courses_workflow():
    """
    Step 9 API End-to-End Test:
    upload demo_courses.csv -> generate -> fetch all 5 views -> export CSV and PDF -> submit 1 feedback row -> all succeed.
    """
    import os

    repo_root = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    demo_csv_path = os.path.join(repo_root, "data", "demo_courses.csv")
    with open(demo_csv_path, "rb") as f:
        csv_bytes = f.read()

    # 1. Upload demo_courses.csv -> generate
    up_resp = client.post(
        "/schedule/upload-csv",
        files={"file": ("demo_courses.csv", csv_bytes, "text/csv")},
    )
    assert up_resp.status_code in (200, 201)
    sched_id = up_resp.json()["schedule_id"]

    # 2. Fetch all 5 views
    views_resp = client.get(f"/schedule/{sched_id}/views")
    assert views_resp.status_code == 200
    views = views_resp.json()["views"]
    for v_name in ("master", "teacher", "student", "staff", "canteen"):
        assert v_name in views
        assert views[v_name]["schedule_id"] == sched_id
        assert "title" in views[v_name]
        assert "generated_at" in views[v_name]

    # 3. Export CSV and PDF
    for fmt in ("csv", "pdf"):
        exp_resp = client.get(f"/export/{sched_id}?format={fmt}&view=teacher")
        assert exp_resp.status_code == 200
        assert len(exp_resp.content) > 100

    # 4. Submit 1 feedback row
    fb_resp = client.post(
        f"/feedback/{sched_id}",
        json={
            "teacher_id": 1,
            "teacher_name": "Dr. Sarah Ahmed",
            "satisfaction_score": 5,
            "preferred_time": "Morning",
            "workload_satisfaction": 5,
            "comments": "Balanced schedule with protected lunch break.",
        },
    )
    assert fb_resp.status_code in (200, 201)
    assert fb_resp.json()["received"] is True


def test_demo_seed_and_shared_fixtures(
    client,
    parsed_demo_schedule,
    solved_demo_entries,
    sample_schedule_payload,
):
    """Verify conftest.py fixtures and POST /demo/seed database fixture seeding."""
    assert len(parsed_demo_schedule["courses"]) == 12
    assert len(solved_demo_entries) == 35
    assert "faculty" in sample_schedule_payload

    seed_resp = client.post("/demo/seed")
    assert seed_resp.status_code == 200
    seed_data = seed_resp.json()
    assert seed_data["schedule_id"] >= 1


